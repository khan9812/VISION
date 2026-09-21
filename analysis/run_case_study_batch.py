"""Run the four GT Case Study images through the VISION analysis flow.

The 006, 026, 066, and 076 cases include SAM segmentation, binary-GT
comparison, two-class CLIP shape classification, and VISION spatial-uniformity
quantification.

The implementation deliberately follows the VISION defaults: BM3D (sigma=40),
Noise2SR (1500 epochs), SAM2 Hiera-L, background-mask removal, and centroid
containment overlap filtering. Results are written per image and aggregated into
CSV and XLSX workbooks. Completed stages are reused on the next invocation.
"""

from __future__ import annotations

import argparse
import gc
import hashlib
import json
import re
import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import cv2
import matplotlib
import numpy as np
import pandas as pd
from scipy.optimize import linear_sum_assignment

matplotlib.use("Agg")
import matplotlib.pyplot as plt


PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))


def configure_console_output() -> None:
    """Avoid aborting a batch when a legacy Windows console cannot encode icons."""
    for stream in (sys.stdout, sys.stderr):
        reconfigure = getattr(stream, "reconfigure", None)
        if reconfigure is None:
            continue
        try:
            reconfigure(errors="backslashreplace")
        except (OSError, ValueError):
            pass


configure_console_output()


from modules.preprocessing import apply_clahe, auto_preprocess
from modules.distribution_analysis import (
    HDBSCAN_AVAILABLE,
    calculate_distribution_metrics,
    compute_unified_voronoi_areas,
    create_clustering_map,
    create_distribution_violin_plot,
    create_voronoi_overlay,
    perform_hdbscan_clustering,
)
from modules.sam2_utils import load_sam_model, filter_overlapping_masks_by_centroid, get_largest_contour_centroid
from modules.runtime_config import resolve_noise2sr_config, runtime_signature
from modules.scientific_plotting import (
    SPATIAL_BOOTSTRAP_RESAMPLES,
    SPATIAL_BOOTSTRAP_SEED,
    SPATIAL_KDE_GRID_POINTS,
    SPATIAL_KDE_HEIGHT,
    SPATIAL_KDE_X_PADDING_FRACTION,
    spatial_distribution_plot_data,
    summarize_spatial_distribution,
)
from modules.shape_analysis import (
    calculate_shape_metrics,
    classify_shapes_with_clip,
    create_shape_overlay,
    create_shape_pie_chart,
)
from modules.visualization import (
    filter_masks_based_on_centroid_containment,
    find_optimal_alpha,
)


PIPELINE_VERSION = "case-study-vision-v2"
SPATIAL_PIPELINE_VERSION = "case-study-spatial-v6"
GT_ORIGIN_WORKBOOK_NAME = "gt_origin_plot_data.xlsx"
GT_IMAGE_STEMS = ("006", "026", "066", "076")
FIGURE_ASSET_FILENAMES = {
    "sam": "sam_segmentation_notext.png",
    "shape": "shape_classification_notext.png",
    "voronoi": "voronoi_tessellation_notext.png",
}
SAM_MASK_COLOR_RGB = (42, 157, 143)
GT_DIST_PLOT_COLUMNS = (
    "Observed_X",
    "Observed_Y",
    "KDE_X",
    "KDE_Y_Normalized",
    "KDE_Y_Plot",
    "KDE_Base_Y",
    "IQR_X",
    "IQR_Y",
    "Median_CI_X",
    "Median_CI_Y",
    "Median_X",
    "Median_Y",
)
VISION_SAM_CONFIG = {
    "model_type": "hiera_l",
    "points_per_side": 32,
    "points_per_batch": 256,
    "pred_iou_thresh": 0.95,
    "stability_score_thresh": 0.80,
    "crop_n_layers": 1,
    "crop_n_points_downscale_factor": 2,
    "crop_nms_thresh": 0.7,
    "box_nms_thresh": 0.7,
    "use_m2m": True,
}

GT_SHAPE_LABELS = ["particle aggregate", "individual particle"]
GT_SHAPE_DESCRIPTIONS = {
    "particle aggregate": "a nanoparticle aggregate",
    "individual particle": "an individual nanoparticle",
}
GT_SHAPE_COLORS_RGB = {
    "individual particle": (42 / 255.0, 157 / 255.0, 143 / 255.0),
    "particle aggregate": (231 / 255.0, 111 / 255.0, 81 / 255.0),
}

@dataclass(frozen=True)
class InputSpec:
    """One scheduled Case Study image."""

    group: str
    image_path: Path
    gt_path: Optional[Path] = None

    @property
    def image_id(self) -> str:
        stem = re.sub(r"[^A-Za-z0-9._-]+", "_", self.image_path.stem).strip("_")
        return stem or "image"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Batch VISION-style analysis for Case Study images."
    )
    execution_scope = parser.add_mutually_exclusive_group()
    parser.add_argument(
        "--case-study-dir",
        type=Path,
        default=PROJECT_ROOT / "Case Study",
        help="Folder containing the four GT Case Study images and binary masks.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=PROJECT_ROOT / "workspace_outputs" / "case_study" / "vision_bm3d_noise2sr",
        help="Root folder for reusable preprocessing and all outputs.",
    )
    parser.add_argument(
        "--bm3d-sigma",
        type=float,
        default=40.0,
        help="BM3D sigma_psd. VISION default is 40.",
    )
    parser.add_argument(
        "--noise2sr-epochs",
        type=int,
        default=1500,
        help="Noise2SR epochs. VISION default is 1500.",
    )
    parser.add_argument(
        "--match-iou",
        type=float,
        default=0.50,
        help="IoU threshold for GT object matching.",
    )
    parser.add_argument(
        "--gt-min-area",
        type=int,
        default=10,
        help="Ignore GT connected components smaller than this pixel area.",
    )
    parser.add_argument(
        "--shape-batch-size",
        type=int,
        default=64,
        help="CLIP batch size for GT shape classification.",
    )
    parser.add_argument(
        "--hdbscan-min-cluster-size",
        type=int,
        default=5,
        help="Minimum particles per HDBSCAN cluster for GT spatial analysis.",
    )
    parser.add_argument(
        "--hdbscan-min-samples",
        type=int,
        default=3,
        help="HDBSCAN core-neighborhood size for GT spatial analysis.",
    )
    parser.add_argument(
        "--clip-model",
        default="ViT-L/14@336px",
        help="CLIP model used by the VISION shape step.",
    )
    parser.add_argument(
        "--enable-clahe",
        action="store_true",
        help="Optionally add CLAHE after BM3D + Noise2SR. Off by default, as in VISION.",
    )
    parser.add_argument(
        "--clahe-clip-limit",
        type=float,
        default=2.0,
        help="CLAHE clip limit when --enable-clahe is set.",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Re-run preprocessing, segmentation, and all requested analysis stages.",
    )
    parser.add_argument(
        "--force-reprocess",
        action="store_true",
        help="Ignore only the BM3D + Noise2SR preprocessing cache.",
    )
    parser.add_argument(
        "--no-resume",
        action="store_true",
        help="Ignore completed segmentation and analysis stage markers.",
    )
    execution_scope.add_argument(
        "--gt-origin-export-only",
        action="store_true",
        help="Build only the Origin workbook for the four completed GT-image analyses.",
    )
    execution_scope.add_argument(
        "--figure-assets-only",
        action="store_true",
        help=(
            "Build text-free SAM, shape, and Voronoi PNG assets from completed "
            "GT analysis caches without running any model."
        ),
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        help="Validate the Case Study layout and print the planned work without running models.",
    )
    return parser.parse_args()

def now_string() -> str:
    return datetime.now().astimezone().isoformat(timespec="seconds")


def json_ready(value: Any) -> Any:
    """Convert NumPy and pathlib values into JSON-compatible objects."""
    if isinstance(value, Path):
        return str(value)
    if isinstance(value, np.generic):
        return value.item()
    if isinstance(value, np.ndarray):
        return value.tolist()
    if isinstance(value, Mapping):
        return {str(key): json_ready(item) for key, item in value.items()}
    if isinstance(value, (list, tuple, set)):
        return [json_ready(item) for item in value]
    return value


def write_json(path: Path, payload: Mapping[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(json_ready(payload), handle, indent=2, ensure_ascii=False)


def read_json(path: Path) -> Optional[Dict[str, Any]]:
    if not path.exists():
        return None
    try:
        with path.open("r", encoding="utf-8") as handle:
            value = json.load(handle)
        return value if isinstance(value, dict) else None
    except (OSError, json.JSONDecodeError):
        return None


def stable_hash(payload: Mapping[str, Any]) -> str:
    text = json.dumps(json_ready(payload), ensure_ascii=True, sort_keys=True)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()[:16]


def file_fingerprint(path: Path) -> Dict[str, Any]:
    stat = path.stat()
    return {
        "path": str(path.resolve()),
        "size_bytes": stat.st_size,
        "mtime_ns": stat.st_mtime_ns,
    }


def segmentation_source_fingerprint(spec: InputSpec) -> Dict[str, Any]:
    """Fingerprint every input that affects a segmentation/GT result."""
    return {
        "image": file_fingerprint(spec.image_path),
        "ground_truth": file_fingerprint(spec.gt_path) if spec.gt_path is not None else None,
    }


def write_bgr_image(path: Path, image: np.ndarray) -> None:
    """Write an OpenCV BGR or grayscale image, including Unicode-safe paths."""
    path.parent.mkdir(parents=True, exist_ok=True)
    suffix = path.suffix if path.suffix else ".png"
    success, encoded = cv2.imencode(suffix, image)
    if not success:
        raise IOError(f"Failed to encode image: {path}")
    encoded.tofile(str(path))


def write_rgb_image(path: Path, image: np.ndarray) -> None:
    """Write an RGB image through OpenCV's BGR encoder."""
    if image.ndim == 3 and image.shape[2] == 3:
        write_bgr_image(path, cv2.cvtColor(image, cv2.COLOR_RGB2BGR))
    else:
        write_bgr_image(path, image)


def read_image_unchanged(path: Path) -> np.ndarray:
    """Read an image with Unicode-safe file handling."""
    encoded = np.fromfile(str(path), dtype=np.uint8)
    image = cv2.imdecode(encoded, cv2.IMREAD_UNCHANGED)
    if image is None:
        raise ValueError(f"Could not load image: {path}")
    return image


def to_bgr_uint8(image: np.ndarray) -> np.ndarray:
    """Normalize TIFF/PNG inputs into contiguous uint8 BGR for VISION modules."""
    if image.ndim == 2:
        normalized = image
        if normalized.dtype != np.uint8:
            normalized = cv2.normalize(normalized, None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
        return np.ascontiguousarray(cv2.cvtColor(normalized, cv2.COLOR_GRAY2BGR))

    if image.ndim != 3:
        raise ValueError(f"Unsupported image shape: {image.shape}")

    if image.shape[2] == 1:
        return to_bgr_uint8(image[:, :, 0])
    if image.shape[2] == 4:
        image = cv2.cvtColor(image, cv2.COLOR_BGRA2BGR)
    elif image.shape[2] != 3:
        raise ValueError(f"Unsupported image channels: {image.shape}")

    if image.dtype != np.uint8:
        image = cv2.normalize(image, None, 0, 255, cv2.NORM_MINMAX).astype(np.uint8)
    return np.ascontiguousarray(image)


def read_bgr_uint8(path: Path) -> np.ndarray:
    return to_bgr_uint8(read_image_unchanged(path))


def image_dir(output_dir: Path, spec: InputSpec) -> Path:
    if spec.group != "gt_evaluation":
        raise ValueError(f"Unsupported Case Study group: {spec.group}")
    return output_dir / "gt_evaluation" / spec.image_id


def preprocessing_cache_paths(output_dir: Path, spec: InputSpec) -> Tuple[Path, Path]:
    cache_dir = output_dir / "preprocessed_bm3d_noise2sr"
    return (
        cache_dir / f"{spec.image_id}_bm3d_noise2sr.png",
        cache_dir / f"{spec.image_id}_bm3d_noise2sr.json",
    )


def cleanup_gpu() -> None:
    """Release temporary tensors between expensive preprocessing/model stages."""
    gc.collect()
    try:
        import torch

        if torch.cuda.is_available():
            torch.cuda.empty_cache()
    except Exception:
        pass


def build_plan(case_study_dir: Path) -> List[InputSpec]:
    return [
        InputSpec(
            group="gt_evaluation",
            image_path=case_study_dir / f"{stem}.tif",
            gt_path=case_study_dir / f"{stem}_gt.tif",
        )
        for stem in GT_IMAGE_STEMS
    ]


def validate_plan(plan: Sequence[InputSpec]) -> List[Path]:
    missing: List[Path] = []
    for spec in plan:
        if not spec.image_path.exists():
            missing.append(spec.image_path)
        if spec.gt_path is not None and not spec.gt_path.exists():
            missing.append(spec.gt_path)
    return missing


def preprocessing_config(args: argparse.Namespace) -> Dict[str, Any]:
    return {
        "pipeline_version": PIPELINE_VERSION,
        "runtime": runtime_signature(),
        "noise2sr": resolve_noise2sr_config(
            {'num_workers': 4, 'persistent_workers': True, 'loader_generator_seed': 42},
            epochs=args.noise2sr_epochs,
        ),
        "method": "bm3d+noise2sr" + ("+clahe" if args.enable_clahe else ""),
        "bm3d_sigma": args.bm3d_sigma,
        "noise2sr_epochs": args.noise2sr_epochs,
        "clahe_enabled": args.enable_clahe,
        "clahe_clip_limit": args.clahe_clip_limit if args.enable_clahe else None,
    }


def segmentation_config(sam_config: Optional[Mapping[str, Any]] = None) -> Dict[str, Any]:
    return {
        "pipeline_version": PIPELINE_VERSION,
        "sam": dict(sam_config or VISION_SAM_CONFIG),
        "background_removal": {
            "bbox_coverage_max": 0.90,
            "border_tolerance": 5,
        },
        "overlap_filter": "centroid_containment",
    }


def shape_config(args: argparse.Namespace) -> Dict[str, Any]:
    return {
        "pipeline_version": PIPELINE_VERSION,
        "clip_model": args.clip_model,
        "shape_labels": GT_SHAPE_LABELS,
        "shape_descriptions": GT_SHAPE_DESCRIPTIONS,
        "clip_batch_size": args.shape_batch_size,
        "confidence_threshold": None,
    }


def spatial_config(args: argparse.Namespace) -> Dict[str, Any]:
    return {
        "pipeline_version": SPATIAL_PIPELINE_VERSION,
        "source": "final_sam_masks_from_bm3d_noise2sr_image",
        "measurement_unit": "px",
        "voronoi_method": "boundary_pixels_clipped_to_convex_hull",
        "kde_kernel": "gaussian",
        "kde_grid_points": SPATIAL_KDE_GRID_POINTS,
        "kde_x_padding_fraction": SPATIAL_KDE_X_PADDING_FRACTION,
        "kde_gui_height": SPATIAL_KDE_HEIGHT,
        "median_ci_method": "percentile_bootstrap",
        "median_ci_confidence": 0.95,
        "median_ci_bootstrap_resamples": SPATIAL_BOOTSTRAP_RESAMPLES,
        "median_ci_seed": SPATIAL_BOOTSTRAP_SEED,
        "hdbscan_min_cluster_size": args.hdbscan_min_cluster_size,
        "hdbscan_min_samples": args.hdbscan_min_samples,
    }


def cache_is_valid(
    metadata: Optional[Mapping[str, Any]],
    config_hash: str,
    fingerprint: Mapping[str, Any],
    required_paths: Iterable[Path],
) -> bool:
    if not metadata or metadata.get("status") != "complete":
        return False
    if metadata.get("config_hash") != config_hash:
        return False
    if metadata.get("source_fingerprint") != json_ready(fingerprint):
        return False
    return all(path.exists() for path in required_paths)


def load_or_preprocess(
    spec: InputSpec,
    args: argparse.Namespace,
    preproc_hash: str,
) -> Tuple[np.ndarray, Dict[str, Any], Path, bool]:
    """Reuse or generate an auditable BM3D + Noise2SR preprocessing result."""
    output_dir = args.output_dir
    image_cache_path, metadata_path = preprocessing_cache_paths(output_dir, spec)
    fingerprint = file_fingerprint(spec.image_path)
    cached_metadata = read_json(metadata_path)
    can_reuse = (
        not args.force
        and not args.force_reprocess
        and not args.no_resume
        and cache_is_valid(cached_metadata, preproc_hash, fingerprint, [image_cache_path])
    )

    if can_reuse:
        cached = read_bgr_uint8(image_cache_path)
        print(f"  [REUSE] Preprocessing: {spec.image_id}")
        return cached, dict(cached_metadata), image_cache_path, True

    source = read_bgr_uint8(spec.image_path)
    print(f"  [RUN] BM3D + Noise2SR preprocessing: {spec.image_id}")
    processed, info = auto_preprocess(
        source,
        sigma_psd=args.bm3d_sigma,
        force_noise2sr=True,
        noise2sr_epochs=args.noise2sr_epochs,
        noise2sr_settings=preprocessing_config(args)['noise2sr'],
    )
    if not (info.get("bm3d_applied") and info.get("noise2sr_applied")):
        raise RuntimeError(
            "VISION preprocessing did not complete BM3D + Noise2SR: "
            f"{info.get('denoising_method', 'unknown')}"
        )

    processed = to_bgr_uint8(processed)
    if args.enable_clahe:
        processed = apply_clahe(processed, args.clahe_clip_limit, (8, 8))
        info["clahe_applied"] = True
        info["clahe_params"] = {"clip_limit": args.clahe_clip_limit, "tile_grid_size": [8, 8]}
    else:
        info["clahe_applied"] = False

    write_bgr_image(image_cache_path, processed)
    metadata = {
        "status": "complete",
        "created_at": now_string(),
        "config_hash": preproc_hash,
        "source_fingerprint": fingerprint,
        "source_shape": list(source.shape),
        "processed_shape": list(processed.shape),
        "preprocessing": info,
        "cache_image": str(image_cache_path),
    }
    write_json(metadata_path, metadata)
    cleanup_gpu()
    return processed, metadata, image_cache_path, False


def mask_bbox(mask: np.ndarray) -> Tuple[int, int, int, int]:
    binary = mask.astype(np.uint8)
    x, y, width, height = cv2.boundingRect(binary)
    return int(x), int(y), int(width), int(height)


def centroid_for_mask(mask: np.ndarray) -> Tuple[float, float]:
    ys, xs = np.where(mask)
    if len(xs) == 0:
        return 0.0, 0.0
    return float(np.mean(xs)), float(np.mean(ys))


def filter_background_masks(masks: Sequence[Dict[str, Any]], image_shape: Tuple[int, int]) -> Tuple[List[Dict[str, Any]], List[Dict[str, Any]]]:
    """Apply the VISION background-mask rule: >90% coverage or four-border touch."""
    height, width = image_shape[:2]
    kept: List[Dict[str, Any]] = []
    background: List[Dict[str, Any]] = []
    border_tolerance = 5

    for mask_dict in masks:
        mask = mask_dict["segmentation"].astype(bool)
        bbox = tuple(mask_dict.get("bbox") or mask_bbox(mask))
        x, y, bbox_width, bbox_height = [int(value) for value in bbox]
        coverage = (bbox_width * bbox_height) / float(height * width)
        touches_all_borders = (
            x <= border_tolerance
            and y <= border_tolerance
            and (x + bbox_width) >= (width - border_tolerance)
            and (y + bbox_height) >= (height - border_tolerance)
        )
        if coverage > 0.90 or touches_all_borders:
            background.append(mask_dict)
        else:
            kept.append(mask_dict)
    return kept, background


def build_sam_mask_generator(model: Any, sam_config: Mapping[str, Any]):
    from sam2.automatic_mask_generator import SAM2AutomaticMaskGenerator

    return SAM2AutomaticMaskGenerator(
        model=model,
        points_per_side=sam_config["points_per_side"],
        points_per_batch=sam_config["points_per_batch"],
        pred_iou_thresh=sam_config["pred_iou_thresh"],
        stability_score_thresh=sam_config["stability_score_thresh"],
        crop_n_layers=sam_config["crop_n_layers"],
        crop_n_points_downscale_factor=sam_config["crop_n_points_downscale_factor"],
        crop_nms_thresh=sam_config["crop_nms_thresh"],
        box_nms_thresh=sam_config["box_nms_thresh"],
        use_m2m=sam_config["use_m2m"],
    )


def generate_sam_masks(
    model: Any,
    image_bgr: np.ndarray,
    sam_config: Mapping[str, Any],
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Generate masks once, with the selected VISION parameters and fallback."""
    sam_input = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB)
    try:
        generator = build_sam_mask_generator(model, sam_config)
        masks = generator.generate(sam_input)
        return masks, {"crop_n_layers_used": sam_config["crop_n_layers"], "fallback_used": False}
    except IndexError as exc:
        if "too many indices for tensor of dimension 1" not in str(exc):
            raise
        fallback_config = dict(sam_config)
        fallback_config["crop_n_layers"] = 0
        print("  [WARN] SAM2 crop-layer empty-result edge case. Retrying with crop_n_layers=0.")
        generator = build_sam_mask_generator(model, fallback_config)
        masks = generator.generate(sam_input)
        return masks, {"crop_n_layers_used": 0, "fallback_used": True}


def run_vision_segmentation(
    model: Any,
    image_bgr: np.ndarray,
    sam_config: Mapping[str, Any],
) -> Tuple[List[Dict[str, Any]], Dict[str, Any]]:
    """Mirror VISION's SAM generation, background removal, and overlap filtering."""
    raw_masks, sam_run_info = generate_sam_masks(model, image_bgr, sam_config)
    normalized_masks: List[Dict[str, Any]] = []
    for mask_dict in raw_masks:
        segmentation = mask_dict.get("segmentation")
        if segmentation is None or not np.any(segmentation):
            continue
        mask = dict(mask_dict)
        mask["segmentation"] = segmentation.astype(bool)
        mask["bbox"] = list(mask.get("bbox") or mask_bbox(mask["segmentation"]))
        mask["area"] = int(mask.get("area", int(mask["segmentation"].sum())))
        normalized_masks.append(mask)

    background_filtered, background_masks = filter_background_masks(
        normalized_masks, image_bgr.shape[:2]
    )
    filtered_masks = filter_overlapping_masks_by_centroid(background_filtered)
    filtered_masks = [mask for mask in filtered_masks
                      if get_largest_contour_centroid(mask['segmentation']) is not None]
    filtered_masks = sorted(
        filtered_masks,
        key=lambda item: (mask_bbox(item["segmentation"])[1], mask_bbox(item["segmentation"])[0]),
    )

    info = {
        "raw_mask_count": len(raw_masks),
        "background_removed_count": len(background_masks),
        "after_background_count": len(background_filtered),
        "final_mask_count": len(filtered_masks),
        "sam_run": sam_run_info,
    }
    return filtered_masks, info


def mask_properties(masks: Sequence[Mapping[str, Any]]) -> List[Dict[str, Any]]:
    properties: List[Dict[str, Any]] = []
    for index, mask_dict in enumerate(masks, start=1):
        mask = mask_dict["segmentation"].astype(bool)
        x, y, width, height = mask_bbox(mask)
        cx, cy = centroid_for_mask(mask)
        properties.append(
            {
                "particle_id": index,
                "area_px": int(mask.sum()),
                "bbox_x": x,
                "bbox_y": y,
                "bbox_width": width,
                "bbox_height": height,
                "centroid_x": round(cx, 3),
                "centroid_y": round(cy, 3),
                "predicted_iou": float(mask_dict.get("predicted_iou", np.nan)),
                "stability_score": float(mask_dict.get("stability_score", np.nan)),
            }
        )
    return properties


def save_masks(path: Path, masks: Sequence[Mapping[str, Any]], image_shape: Tuple[int, int]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if masks:
        array = np.stack([mask["segmentation"].astype(np.uint8) for mask in masks], axis=0)
    else:
        array = np.empty((0, image_shape[0], image_shape[1]), dtype=np.uint8)
    np.savez_compressed(path, masks=array)


def load_masks(path: Path) -> List[Dict[str, Any]]:
    with np.load(path) as data:
        array = data["masks"].astype(bool)
    masks: List[Dict[str, Any]] = []
    for mask in array:
        masks.append(
            {
                "segmentation": mask,
                "bbox": list(mask_bbox(mask)),
                "area": int(mask.sum()),
            }
        )
    return masks


def mask_union(masks: Sequence[Mapping[str, Any]], shape: Tuple[int, int]) -> np.ndarray:
    union = np.zeros(shape, dtype=bool)
    for mask in masks:
        union |= mask["segmentation"].astype(bool)
    return union


def segmentation_overlay(image_bgr: np.ndarray, masks: Sequence[Mapping[str, Any]]) -> np.ndarray:
    """Render final VISION-style particle masks with white outlines."""
    base = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB).astype(np.float32)
    palette = [
        (79, 124, 172),
        (42, 157, 143),
        (217, 109, 59),
        (123, 97, 168),
        (118, 183, 178),
        (233, 196, 106),
    ]
    for index, mask_dict in enumerate(sorted(masks, key=lambda item: item["area"], reverse=True)):
        mask = mask_dict["segmentation"].astype(bool)
        color = np.asarray(palette[index % len(palette)], dtype=np.float32)
        base[mask] = base[mask] * 0.48 + color * 0.52
        contours, _ = cv2.findContours(mask.astype(np.uint8), cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        cv2.drawContours(base, contours, -1, (255, 255, 255), 1, lineType=cv2.LINE_AA)
    return np.clip(base, 0, 255).astype(np.uint8)


def class_colored_mask_overlay(
    image_bgr: np.ndarray,
    masks: Sequence[Mapping[str, Any]],
    colors_rgb: Sequence[Sequence[float]],
    alpha: float = 0.52,
) -> np.ndarray:
    """Render a text-free, source-resolution mask overlay with white outlines."""
    if len(masks) != len(colors_rgb):
        raise ValueError(
            f"Mask/color count mismatch: {len(masks)} masks vs {len(colors_rgb)} colors"
        )

    overlay = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB).astype(np.float32)
    contours_by_mask = []
    for mask_dict, raw_color in zip(masks, colors_rgb):
        mask = np.asarray(mask_dict["segmentation"], dtype=bool)
        if mask.shape != image_bgr.shape[:2]:
            raise ValueError(
                f"Mask/image shape mismatch: {mask.shape} vs {image_bgr.shape[:2]}"
            )
        color = np.asarray(raw_color, dtype=np.float32)
        if color.shape != (3,):
            raise ValueError(f"Expected an RGB triplet, got: {raw_color}")
        if float(np.max(color)) <= 1.0:
            color = color * 255.0
        overlay[mask] = overlay[mask] * (1.0 - alpha) + color * alpha
        contours, _ = cv2.findContours(
            mask.astype(np.uint8), cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE
        )
        contours_by_mask.append(contours)

    for contours in contours_by_mask:
        cv2.drawContours(
            overlay, contours, -1, (255, 255, 255), 1, lineType=cv2.LINE_AA
        )
    return np.clip(overlay, 0, 255).astype(np.uint8)


def sam_mask_overlay_notext(
    image_bgr: np.ndarray,
    masks: Sequence[Mapping[str, Any]],
) -> np.ndarray:
    """Render all accepted SAM masks in one consistent publication color."""
    return class_colored_mask_overlay(
        image_bgr,
        masks,
        [SAM_MASK_COLOR_RGB] * len(masks),
    )


def shape_mask_overlay_notext(
    image_bgr: np.ndarray,
    masks: Sequence[Mapping[str, Any]],
    shapes: Sequence[str],
    color_map: Mapping[str, Sequence[float]],
) -> np.ndarray:
    """Render one stable color per classified shape without labels or legend."""
    if len(masks) != len(shapes):
        raise ValueError(
            f"Mask/shape count mismatch: {len(masks)} masks vs {len(shapes)} labels"
        )
    unexpected = sorted(set(shapes) - set(color_map))
    if unexpected:
        raise ValueError(f"No display color is defined for shapes: {unexpected}")
    return class_colored_mask_overlay(
        image_bgr,
        masks,
        [color_map[shape] for shape in shapes],
    )


def save_segmentation_visuals(
    destination: Path,
    image_bgr: np.ndarray,
    masks: Sequence[Mapping[str, Any]],
) -> np.ndarray:
    union = mask_union(masks, image_bgr.shape[:2])
    overlay = segmentation_overlay(image_bgr, masks)
    source_rgb = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB)
    particles_only = source_rgb.copy()
    particles_only[~union] = 255

    write_rgb_image(destination / "segmentation_overlay.png", overlay)
    write_rgb_image(
        destination / FIGURE_ASSET_FILENAMES["sam"],
        sam_mask_overlay_notext(image_bgr, masks),
    )
    write_rgb_image(destination / "background_removed_particles.png", particles_only)
    write_bgr_image(destination / "sam_background_mask.png", np.where(union, 0, 255).astype(np.uint8))
    write_bgr_image(destination / "predicted_union_mask.png", np.where(union, 255, 0).astype(np.uint8))
    return union


def binary_gt_instances(gt_path: Path, min_area: int) -> Tuple[List[np.ndarray], List[Dict[str, Any]], np.ndarray]:
    image = read_image_unchanged(gt_path)
    if image.ndim == 3:
        if image.shape[2] == 4:
            image = cv2.cvtColor(image, cv2.COLOR_BGRA2GRAY)
        else:
            image = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    binary = image > 0
    count, labels, stats, centroids = cv2.connectedComponentsWithStats(binary.astype(np.uint8), connectivity=8)

    instances: List[np.ndarray] = []
    properties: List[Dict[str, Any]] = []
    for component_id in range(1, count):
        area = int(stats[component_id, cv2.CC_STAT_AREA])
        if area < min_area:
            continue
        instance = labels == component_id
        instances.append(instance)
        properties.append(
            {
                "gt_id": len(instances),
                "area_px": area,
                "bbox_x": int(stats[component_id, cv2.CC_STAT_LEFT]),
                "bbox_y": int(stats[component_id, cv2.CC_STAT_TOP]),
                "bbox_width": int(stats[component_id, cv2.CC_STAT_WIDTH]),
                "bbox_height": int(stats[component_id, cv2.CC_STAT_HEIGHT]),
                "centroid_x": round(float(centroids[component_id, 0]), 3),
                "centroid_y": round(float(centroids[component_id, 1]), 3),
            }
        )
    clean_union = np.zeros(binary.shape, dtype=bool)
    for instance in instances:
        clean_union |= instance
    return instances, properties, clean_union


def mask_iou(mask_a: np.ndarray, mask_b: np.ndarray) -> float:
    intersection = np.logical_and(mask_a, mask_b).sum()
    union = np.logical_or(mask_a, mask_b).sum()
    return float(intersection / union) if union else 0.0


def evaluate_against_gt(
    gt_masks: Sequence[np.ndarray],
    prediction_masks: Sequence[Mapping[str, Any]],
    iou_threshold: float,
) -> Tuple[Dict[str, Any], pd.DataFrame, np.ndarray, np.ndarray, np.ndarray]:
    pred_masks = [mask["segmentation"].astype(bool) for mask in prediction_masks]
    if gt_masks:
        shape = gt_masks[0].shape
    elif pred_masks:
        shape = pred_masks[0].shape
    else:
        raise ValueError("GT and prediction masks are both empty")
    if any(mask.shape != shape for mask in gt_masks) or any(mask.shape != shape for mask in pred_masks):
        raise ValueError("GT and predicted masks have different image shapes")

    iou_matrix = np.zeros((len(gt_masks), len(pred_masks)), dtype=float)
    for gt_index, gt_mask in enumerate(gt_masks):
        for pred_index, pred_mask in enumerate(pred_masks):
            iou_matrix[gt_index, pred_index] = mask_iou(gt_mask, pred_mask)

    matches: List[Tuple[int, int, float]] = []
    if len(gt_masks) and len(pred_masks):
        rows, columns = linear_sum_assignment(-iou_matrix)
        for gt_index, pred_index in zip(rows, columns):
            iou = float(iou_matrix[gt_index, pred_index])
            if iou >= iou_threshold:
                matches.append((int(gt_index), int(pred_index), iou))

    matched_gt = {item[0] for item in matches}
    matched_pred = {item[1] for item in matches}
    object_precision = len(matches) / len(pred_masks) if pred_masks else 0.0
    object_recall = len(matches) / len(gt_masks) if gt_masks else 0.0
    object_f1 = (
        2 * object_precision * object_recall / (object_precision + object_recall)
        if (object_precision + object_recall)
        else 0.0
    )

    gt_union = np.zeros(shape, dtype=bool)
    pred_union = np.zeros(shape, dtype=bool)
    for mask in gt_masks:
        gt_union |= mask
    for mask in pred_masks:
        pred_union |= mask
    true_positive = gt_union & pred_union
    false_positive = pred_union & ~gt_union
    false_negative = gt_union & ~pred_union
    pixel_precision = true_positive.sum() / pred_union.sum() if pred_union.any() else 0.0
    pixel_recall = true_positive.sum() / gt_union.sum() if gt_union.any() else 0.0
    pixel_f1 = (
        2 * pixel_precision * pixel_recall / (pixel_precision + pixel_recall)
        if (pixel_precision + pixel_recall)
        else 0.0
    )

    match_rows: List[Dict[str, Any]] = [
        {"status": "matched", "gt_id": gt_index + 1, "predicted_id": pred_index + 1, "iou": iou}
        for gt_index, pred_index, iou in matches
    ]
    match_rows.extend(
        {"status": "false_negative", "gt_id": gt_index + 1, "predicted_id": None, "iou": 0.0}
        for gt_index in range(len(gt_masks))
        if gt_index not in matched_gt
    )
    match_rows.extend(
        {"status": "false_positive", "gt_id": None, "predicted_id": pred_index + 1, "iou": 0.0}
        for pred_index in range(len(pred_masks))
        if pred_index not in matched_pred
    )

    metrics = {
        "match_iou_threshold": iou_threshold,
        "gt_instance_count": len(gt_masks),
        "predicted_instance_count": len(pred_masks),
        "matched_instance_count": len(matches),
        "false_positive_count": len(pred_masks) - len(matches),
        "false_negative_count": len(gt_masks) - len(matches),
        "object_precision": object_precision,
        "object_recall": object_recall,
        "object_f1": object_f1,
        "mean_matched_iou": float(np.mean([item[2] for item in matches])) if matches else 0.0,
        "median_matched_iou": float(np.median([item[2] for item in matches])) if matches else 0.0,
        "pixel_true_positive": int(true_positive.sum()),
        "pixel_false_positive": int(false_positive.sum()),
        "pixel_false_negative": int(false_negative.sum()),
        "pixel_precision": float(pixel_precision),
        "pixel_recall": float(pixel_recall),
        "pixel_f1": float(pixel_f1),
    }
    return metrics, pd.DataFrame(match_rows), iou_matrix, gt_union, pred_union


def gt_error_overlay(image_bgr: np.ndarray, gt_union: np.ndarray, pred_union: np.ndarray) -> np.ndarray:
    """Green = TP, red = FP, blue = FN over the preprocessed image."""
    base = cv2.cvtColor(image_bgr, cv2.COLOR_BGR2RGB).astype(np.float32)
    layers = [
        (gt_union & pred_union, np.asarray((50, 170, 95), dtype=np.float32)),
        (pred_union & ~gt_union, np.asarray((215, 75, 70), dtype=np.float32)),
        (gt_union & ~pred_union, np.asarray((70, 125, 220), dtype=np.float32)),
    ]
    for mask, color in layers:
        base[mask] = base[mask] * 0.38 + color * 0.62
    return np.clip(base, 0, 255).astype(np.uint8)


def segmentation_summary_path(destination: Path) -> Path:
    return destination / "segmentation_summary.json"


def segmentation_cache_valid(
    spec: InputSpec,
    args: argparse.Namespace,
    segmentation_hash: str,
) -> bool:
    if args.force or args.no_resume:
        return False
    return segmentation_artifacts_valid(spec, args.output_dir, segmentation_hash)


def segmentation_artifacts_valid(
    spec: InputSpec,
    output_dir: Path,
    segmentation_hash: str,
) -> bool:
    """Check completed segmentation artifacts independently of run-control flags."""
    destination = image_dir(output_dir, spec)
    summary = read_json(segmentation_summary_path(destination))
    if not summary:
        return False
    try:
        fingerprint = segmentation_source_fingerprint(spec)
    except FileNotFoundError:
        return False
    return cache_is_valid(
        summary,
        segmentation_hash,
        fingerprint,
        [destination / "predicted_masks.npz", destination / "mask_properties.csv"],
    )


def run_segmentation_for_image(
    model: Any,
    spec: InputSpec,
    args: argparse.Namespace,
    preproc_metadata: Mapping[str, Any],
    preprocessed_path: Path,
    segmentation_hash: str,
    sam_config: Mapping[str, Any],
) -> Dict[str, Any]:
    destination = image_dir(args.output_dir, spec)
    destination.mkdir(parents=True, exist_ok=True)
    processed = read_bgr_uint8(preprocessed_path)
    masks, segmentation_info = run_vision_segmentation(model, processed, sam_config)
    properties = mask_properties(masks)

    write_bgr_image(destination / "preprocessed_bm3d_noise2sr.png", processed)
    write_json(destination / "preprocessing.json", dict(preproc_metadata))
    save_masks(destination / "predicted_masks.npz", masks, processed.shape[:2])
    pd.DataFrame(properties).to_csv(destination / "mask_properties.csv", index=False)
    predicted_union = save_segmentation_visuals(destination, processed, masks)

    summary: Dict[str, Any] = {
        "status": "complete",
        "stage": "segmentation",
        "created_at": now_string(),
        "config_hash": segmentation_hash,
        "source_fingerprint": segmentation_source_fingerprint(spec),
        "image_id": spec.image_id,
        "group": spec.group,
        "source_image": str(spec.image_path),
        "preprocessed_image": str(preprocessed_path),
        "input_shape": list(read_bgr_uint8(spec.image_path).shape),
        "processed_shape": list(processed.shape),
        "sam_config": dict(sam_config),
        "segmentation": segmentation_info,
        "outputs": {
            "masks": "predicted_masks.npz",
            "mask_properties": "mask_properties.csv",
            "overlay": "segmentation_overlay.png",
            "sam_notext": FIGURE_ASSET_FILENAMES["sam"],
            "background_mask": "sam_background_mask.png",
        },
    }

    if spec.gt_path is not None:
        gt_masks, gt_properties, gt_union = binary_gt_instances(spec.gt_path, args.gt_min_area)
        if gt_union.shape != predicted_union.shape:
            raise ValueError(
                f"GT/prediction shape mismatch for {spec.image_id}: "
                f"{gt_union.shape} vs {predicted_union.shape}"
            )
        metrics, matches, iou_matrix, _, _ = evaluate_against_gt(
            gt_masks, masks, args.match_iou
        )
        write_bgr_image(destination / "ground_truth_union_mask.png", np.where(gt_union, 255, 0).astype(np.uint8))
        pd.DataFrame(gt_properties).to_csv(destination / "ground_truth_properties.csv", index=False)
        matches.to_csv(destination / "instance_matches.csv", index=False)
        pd.DataFrame(iou_matrix).to_csv(destination / "iou_matrix.csv", index=False)
        write_rgb_image(destination / "gt_pixel_error_overlay.png", gt_error_overlay(processed, gt_union, predicted_union))
        write_json(destination / "gt_metrics.json", metrics)
        summary["gt_evaluation"] = metrics

    write_json(segmentation_summary_path(destination), summary)
    print(
        f"  [OK] Segmentation {spec.image_id}: "
        f"raw={segmentation_info['raw_mask_count']}, "
        f"background_removed={segmentation_info['background_removed_count']}, "
        f"final={len(masks)}"
    )
    if spec.gt_path is not None:
        evaluation = summary["gt_evaluation"]
        print(
            f"       Object P/R/F1 = {evaluation['object_precision']:.3f} / "
            f"{evaluation['object_recall']:.3f} / {evaluation['object_f1']:.3f}"
        )
    cleanup_gpu()
    return summary


def shape_summary_path(destination: Path) -> Path:
    return destination / "shape_summary.json"


def shape_cache_valid(spec: InputSpec, args: argparse.Namespace, shape_hash: str) -> bool:
    destination = image_dir(args.output_dir, spec)
    summary = read_json(shape_summary_path(destination))
    if args.force or args.no_resume:
        return False
    return cache_is_valid(
        summary,
        shape_hash,
        file_fingerprint(spec.image_path),
        [destination / "particle_shapes.csv"],
    )


def shape_text_descriptions(
    labels: Sequence[str],
    descriptions: Mapping[str, str],
) -> List[str]:
    return [
        f"This is {descriptions[label]} on an electron microscope image"
        for label in labels
    ]


def load_clip_resources(clip_model_name: str) -> Dict[str, Any]:
    import clip
    import torch

    device = "cuda" if torch.cuda.is_available() else "cpu"
    print(f"Loading CLIP {clip_model_name} on {device}...")
    model, preprocess = clip.load(clip_model_name, device=device)
    model.eval()
    return {
        "model": model,
        "preprocess": preprocess,
        "device": device,
    }


def save_figure(fig: plt.Figure, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(path, dpi=300, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def spatial_summary_path(destination: Path) -> Path:
    return destination / "spatial_summary.json"


def spatial_source_fingerprint(spec: InputSpec, output_dir: Path) -> Dict[str, Any]:
    destination = image_dir(output_dir, spec)
    preprocessed_path, _ = preprocessing_cache_paths(output_dir, spec)
    return {
        "image": file_fingerprint(spec.image_path),
        "preprocessed_bm3d_noise2sr": file_fingerprint(preprocessed_path),
        "predicted_masks": file_fingerprint(destination / "predicted_masks.npz"),
    }


def spatial_cache_valid(
    spec: InputSpec,
    args: argparse.Namespace,
    spatial_hash: str,
) -> bool:
    if args.force or args.no_resume:
        return False
    destination = image_dir(args.output_dir, spec)
    try:
        fingerprint = spatial_source_fingerprint(spec, args.output_dir)
    except FileNotFoundError:
        return False
    return cache_is_valid(
        read_json(spatial_summary_path(destination)),
        spatial_hash,
        fingerprint,
        [
            destination / "spatial_metrics.json",
            destination / "spatial_voronoi_areas.csv",
            destination / "spatial_distribution.png",
            destination / "spatial_voronoi_overlay.png",
            destination / "spatial_clustering.png",
        ],
    )


def extract_mask_boundary_records(
    masks: Sequence[Mapping[str, Any]],
) -> List[Dict[str, Any]]:
    """Extract each final SAM mask's contour, centroid, and pixel area once."""
    records: List[Dict[str, Any]] = []
    used_centroids = set()
    for mask_dict in masks:
        binary = mask_dict["segmentation"].astype(np.uint8)
        contours, _ = cv2.findContours(binary, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
        if not contours:
            continue
        largest_contour = max(contours, key=cv2.contourArea)
        moments = cv2.moments(largest_contour)
        if moments["m00"]:
            centroid = (
                float(int(moments["m10"] / moments["m00"])),
                float(int(moments["m01"] / moments["m00"])),
            )
        else:
            centroid = centroid_for_mask(binary.astype(bool))
        while centroid in used_centroids:
            centroid = (centroid[0] + 1e-3, centroid[1] + 1e-3)
        points = [
            tuple(int(value) for value in point[0])
            for contour in contours
            for point in contour
        ]
        if not points:
            continue
        used_centroids.add(centroid)
        records.append(
            {
                "particle_id": len(records) + 1,
                "centroid": centroid,
                "boundary_points": points,
                "area_px": int(binary.sum()),
            }
        )
    return records


def spatial_inputs_from_masks(
    masks: Sequence[Mapping[str, Any]],
) -> Tuple[List[Tuple[float, float]], Dict[Tuple[float, float], List[Tuple[int, int]]]]:
    """Extract the centroid-keyed boundary pixels used by the VISION GUI."""
    records = extract_mask_boundary_records(masks)
    centroids = [record["centroid"] for record in records]
    boundary_pixels = {
        record["centroid"]: record["boundary_points"]
        for record in records
    }
    return centroids, boundary_pixels


class PFUnavailableError(ValueError):
    """A valid segmentation has insufficient geometry for sample-SD PF-SUI."""
    def __init__(self, reason, n_pf=0):
        super().__init__(reason)
        self.n_pf = n_pf


def spatial_geometry_from_masks(
    masks: Sequence[Mapping[str, Any]],
) -> Tuple[
    List[Tuple[float, float]],
    Dict[Tuple[float, float], List[Tuple[int, int]]],
    List[Tuple[float, float]],
    List[Tuple[float, float]],
    Dict[Tuple[float, float], float],
    float,
]:
    """Compute the exact VISION Voronoi geometry shared by analysis and export."""
    centroids, boundary_pixels = spatial_inputs_from_masks(masks)
    if len(centroids) < 3:
        raise PFUnavailableError(
            f"Spatial uniformity requires at least 3 particles; found {len(centroids)}"
        )

    inside_centroids, _, concave_vertex, current_alpha = find_optimal_alpha(centroids)
    analysis_centroids = list(inside_centroids or [])
    if len(concave_vertex) < 3:
        raise PFUnavailableError('PF-SUI unavailable: no nonzero-area convex hull')
    if len(analysis_centroids) < 2:
        raise PFUnavailableError(f'PF-SUI unavailable: {len(analysis_centroids)} interior centroids; at least two required', len(analysis_centroids))

    voronoi_areas, _, _ = compute_unified_voronoi_areas(
        boundary_pixels,
        concave_vertex,
        analysis_centroids,
    )
    if len(voronoi_areas) < 2:
        raise PFUnavailableError('PF-SUI unavailable: fewer than two finite interior regions', len(voronoi_areas))

    return (
        centroids,
        boundary_pixels,
        analysis_centroids,
        concave_vertex,
        voronoi_areas,
        float(current_alpha),
    )


def run_spatial_for_image(
    spec: InputSpec,
    args: argparse.Namespace,
    preprocessed_path: Path,
    spatial_hash: str,
) -> Dict[str, Any]:
    """Quantify spatial uniformity from final SAM masks in pixel units."""
    destination = image_dir(args.output_dir, spec)
    masks = load_masks(destination / "predicted_masks.npz")
    processed = read_bgr_uint8(preprocessed_path)
    try:
        (centroids, boundary_pixels, analysis_centroids, concave_vertex,
         voronoi_areas, current_alpha) = spatial_geometry_from_masks(masks)
    except PFUnavailableError as exc:
        metrics = {'particle_count': len(masks), 'n_pf': exc.n_pf,
                   'spatial_uniformity_index': None, 'voronoi_cv_percent': None}
        document = {'status': 'unavailable', 'reason': str(exc), 'measurement_unit': 'px',
                    'area_unit': 'px^2', 'spatial_metrics': metrics, 'clustering_metrics': {}}
        write_json(destination / 'spatial_metrics.json', document)
        pd.DataFrame(columns=['centroid_x','centroid_y','voronoi_area_px2']).to_csv(
            destination / 'spatial_voronoi_areas.csv', index=False)
        for filename in ('spatial_distribution.png', 'spatial_voronoi_overlay.png',
                         'spatial_clustering.png', 'spatial_clusters.csv', FIGURE_ASSET_FILENAMES['voronoi']):
            stale = destination / filename
            if stale.resolve().parent != destination.resolve():
                raise ValueError('Unexpected spatial artifact path')
            stale.unlink(missing_ok=True)
        summary = dict(document, stage='spatial_uniformity', image_id=spec.image_id,
                       config_hash=spatial_hash, created_at=now_string(),
                       source_fingerprint=spatial_source_fingerprint(spec,args.output_dir),
                       outputs={'metrics':'spatial_metrics.json','voronoi_areas':'spatial_voronoi_areas.csv'})
        write_json(spatial_summary_path(destination), summary)
        return summary

    area_values = [float(value) for value in voronoi_areas.values()]
    distribution_metrics = calculate_distribution_metrics(area_values, scale="px")
    distribution_summary = summarize_spatial_distribution(area_values)

    cluster_centroids = analysis_centroids if len(analysis_centroids) >= 3 else centroids
    effective_min_cluster_size = max(
        2, min(args.hdbscan_min_cluster_size, len(cluster_centroids))
    )
    effective_min_samples = max(
        1, min(args.hdbscan_min_samples, max(1, len(cluster_centroids) - 1))
    )
    cluster_labels, cluster_stats = perform_hdbscan_clustering(
        cluster_centroids,
        min_cluster_size=effective_min_cluster_size,
        min_samples=effective_min_samples,
        random_state=42,
    )

    spatial_metrics = {
        "particle_count": len(masks),
        "boundary_particle_count": len(boundary_pixels),
        "concave_hull_particle_count": len(analysis_centroids),
        "analysis_hull_particle_count": len(analysis_centroids),
        "analysis_boundary_method": "convex_hull_with_5px_inward_centroid_filter",
        "voronoi_cell_count": len(area_values),
        "concave_alpha": float(current_alpha),
        "mean_voronoi_area_px2": distribution_metrics["mean"],
        "std_voronoi_area_px2": distribution_metrics["std"],
        "median_voronoi_area_px2": distribution_summary["median"],
        "median_voronoi_ci_low_px2": distribution_summary["median_ci_low"],
        "median_voronoi_ci_high_px2": distribution_summary["median_ci_high"],
        "q1_voronoi_area_px2": distribution_summary["q1"],
        "q3_voronoi_area_px2": distribution_summary["q3"],
        "voronoi_cv_percent": distribution_metrics["cv"],
        "spatial_uniformity_index": distribution_metrics["sui"],
    }
    clustering_metrics = {
        "hdbscan_available": HDBSCAN_AVAILABLE,
        "requested_min_cluster_size": args.hdbscan_min_cluster_size,
        "requested_min_samples": args.hdbscan_min_samples,
        "effective_min_cluster_size": effective_min_cluster_size,
        "effective_min_samples": effective_min_samples,
        "n_clusters": cluster_stats["n_clusters"],
        "noise_points": cluster_stats["n_noise"],
        "clustered_percentage": cluster_stats["clustered_percentage"],
        "avg_cluster_size": cluster_stats["avg_cluster_size"],
        "cluster_fraction": (
            1.0 - cluster_stats["n_noise"] / len(cluster_labels)
            if len(cluster_labels)
            else 0.0
        ),
    }

    area_rows = [
        {
            "centroid_x": float(centroid[0]),
            "centroid_y": float(centroid[1]),
            "voronoi_area_px2": float(area),
        }
        for centroid, area in voronoi_areas.items()
    ]
    pd.DataFrame(area_rows).to_csv(destination / "spatial_voronoi_areas.csv", index=False)
    cluster_rows = [
        {
            "centroid_x": float(centroid[0]),
            "centroid_y": float(centroid[1]),
            "hdbscan_label": int(label),
            "is_noise": bool(label == -1),
        }
        for centroid, label in zip(cluster_centroids, cluster_labels)
    ]
    pd.DataFrame(cluster_rows).to_csv(destination / "spatial_clusters.csv", index=False)

    spatial_figure = create_distribution_violin_plot(area_values, scale="px", figsize=(8, 5))
    save_figure(spatial_figure, destination / "spatial_distribution.png")
    voronoi_figure, voronoi_overlay = create_voronoi_overlay(
        processed,
        boundary_pixels,
        voronoi_areas,
        scale=None,
        figsize=(10, 10),
        concave_vertex=concave_vertex,
        inside_centroids=analysis_centroids,
        show_area_labels=False,
        show_title=False,
        colorbar_label="Relative Voronoi Cell Area",
        relative_colorbar=True,
    )
    save_figure(voronoi_figure, destination / "spatial_voronoi_overlay.png")
    write_bgr_image(
        destination / FIGURE_ASSET_FILENAMES["voronoi"],
        voronoi_overlay,
    )
    clustering_figure, _ = create_clustering_map(
        processed,
        cluster_centroids,
        cluster_labels,
        figsize=(10, 10),
    )
    save_figure(clustering_figure, destination / "spatial_clustering.png")

    metrics_document = {
        "measurement_unit": "px",
        "area_unit": "px^2",
        "spatial_metrics": spatial_metrics,
        "clustering_metrics": clustering_metrics,
    }
    write_json(destination / "spatial_metrics.json", metrics_document)
    summary = {
        "status": "complete",
        "stage": "spatial_uniformity",
        "created_at": now_string(),
        "config_hash": spatial_hash,
        "source_fingerprint": spatial_source_fingerprint(spec, args.output_dir),
        "image_id": spec.image_id,
        "group": spec.group,
        "source_image": str(spec.image_path),
        "preprocessed_image": str(preprocessed_path),
        **metrics_document,
        "outputs": {
            "metrics": "spatial_metrics.json",
            "voronoi_areas": "spatial_voronoi_areas.csv",
            "clusters": "spatial_clusters.csv",
            "distribution": "spatial_distribution.png",
            "voronoi_overlay": "spatial_voronoi_overlay.png",
            "voronoi_notext": FIGURE_ASSET_FILENAMES["voronoi"],
            "clustering": "spatial_clustering.png",
        },
    }
    write_json(spatial_summary_path(destination), summary)
    print(
        f"  [OK] Spatial {spec.image_id}: cells={len(area_values)}, "
        f"CV={spatial_metrics['voronoi_cv_percent']:.2f}%, "
        f"SUI={spatial_metrics['spatial_uniformity_index']:.4f}"
    )
    return summary


def run_shape_for_image(
    spec: InputSpec,
    args: argparse.Namespace,
    clip_resources: Optional[Mapping[str, Any]],
    shape_hash: str,
) -> Dict[str, Any]:
    destination = image_dir(args.output_dir, spec)
    masks = load_masks(destination / "predicted_masks.npz")
    processed = read_bgr_uint8(destination / "preprocessed_bm3d_noise2sr.png")
    properties = mask_properties(masks)
    labels = list(GT_SHAPE_LABELS)
    descriptions = GT_SHAPE_DESCRIPTIONS
    color_map = dict(GT_SHAPE_COLORS_RGB)

    if masks:
        if clip_resources is None:
            raise RuntimeError("CLIP resources are required when masks are present")
        import clip

        text_descriptions = shape_text_descriptions(labels, descriptions)
        text_tokens = clip.tokenize(text_descriptions).to(clip_resources["device"])
        shapes, confidences, shape_counts = classify_shapes_with_clip(
            masks,
            processed,
            labels,
            text_descriptions,
            clip_resources["model"],
            clip_resources["preprocess"],
            text_tokens,
            clip_resources["device"],
            batch_size=args.shape_batch_size,
            temperature=0.07,
            confidence_threshold=None,
        )
        metrics = calculate_shape_metrics(shapes, confidences)
        overlay = create_shape_overlay(processed, masks, shapes, color_map, figsize=(10, 10))
        pie = create_shape_pie_chart(dict(shape_counts), color_map, figsize=(8, 7))
        save_figure(overlay, destination / "shape_overlay.png")
        save_figure(pie, destination / "shape_composition.png")
        write_rgb_image(
            destination / FIGURE_ASSET_FILENAMES["shape"],
            shape_mask_overlay_notext(processed, masks, shapes, color_map),
        )
    else:
        shapes = []
        confidences = []
        shape_counts = {}
        metrics = {
            "total_particles": 0,
            "shape_counts": {},
            "percentages": {},
            "shannon_entropy": 0.0,
            "simpson_diversity": 0.0,
            "dominant_shape": None,
            "dominant_percentage": 0.0,
            "avg_confidence": 0.0,
        }

    particles = pd.DataFrame(properties)
    if particles.empty:
        particles = pd.DataFrame(columns=["particle_id", "area_px", "bbox_x", "bbox_y", "bbox_width", "bbox_height", "centroid_x", "centroid_y", "predicted_iou", "stability_score"])
    particles["shape"] = shapes
    particles["shape_confidence"] = confidences
    particles.to_csv(destination / "particle_shapes.csv", index=False)

    summary = {
        "status": "complete",
        "stage": "shape_classification",
        "created_at": now_string(),
        "config_hash": shape_hash,
        "source_fingerprint": file_fingerprint(spec.image_path),
        "image_id": spec.image_id,
        "group": spec.group,
        "source_image": str(spec.image_path),
        "particle_count": len(masks),
        "shape_labels": labels,
        "shape_metrics": metrics,
        "shape_counts_all_labels": {label: int(shape_counts.get(label, 0)) for label in labels},
        "outputs": {
            "particles": "particle_shapes.csv",
            "overlay": "shape_overlay.png" if masks else None,
            "shape_notext": FIGURE_ASSET_FILENAMES["shape"] if masks else None,
            "composition": "shape_composition.png" if masks else None,
        },
    }
    write_json(shape_summary_path(destination), summary)
    print(f"  [OK] Shape classification {spec.image_id}: {len(masks)} particles")
    cleanup_gpu()
    return summary


def rgb_to_hex(color: Sequence[float]) -> str:
    values = np.asarray(color, dtype=np.float64)
    if values.shape != (3,):
        raise ValueError(f"Expected an RGB triplet, got: {color}")
    if float(np.max(values)) <= 1.0:
        values = values * 255.0
    values = np.clip(np.rint(values), 0, 255).astype(int)
    return "#" + "".join(f"{value:02X}" for value in values)


def export_gt_figure_assets(
    plan: Sequence[InputSpec],
    output_dir: Path,
) -> Path:
    """Create the three text-free figure panels from completed GT caches."""
    records: List[Dict[str, Any]] = []
    color_map = dict(GT_SHAPE_COLORS_RGB)

    for spec in plan:
        if spec.group != "gt_evaluation":
            continue
        destination = image_dir(output_dir, spec)
        preprocessed_path = destination / "preprocessed_bm3d_noise2sr.png"
        masks_path = destination / "predicted_masks.npz"
        shapes_path = destination / "particle_shapes.csv"
        required_paths = [preprocessed_path, masks_path, shapes_path]
        missing = [str(path) for path in required_paths if not path.exists()]
        if missing:
            raise FileNotFoundError(
                f"Figure assets require completed GT caches for {spec.image_id}: "
                + ", ".join(missing)
            )

        processed = read_bgr_uint8(preprocessed_path)
        masks = load_masks(masks_path)
        particle_shapes = pd.read_csv(shapes_path)
        required_columns = {"particle_id", "shape"}
        missing_columns = required_columns - set(particle_shapes.columns)
        if missing_columns:
            raise ValueError(
                f"{shapes_path} is missing columns: "
                + ", ".join(sorted(missing_columns))
            )
        if len(particle_shapes) != len(masks):
            raise ValueError(
                f"Shape/mask count mismatch for {spec.image_id}: "
                f"{len(particle_shapes)} vs {len(masks)}"
            )

        particle_ids = pd.to_numeric(
            particle_shapes["particle_id"], errors="raise"
        ).astype(int)
        expected_ids = np.arange(1, len(masks) + 1, dtype=int)
        if not np.array_equal(particle_ids.to_numpy(), expected_ids):
            raise ValueError(
                f"particle_shapes.csv order does not match mask order for {spec.image_id}"
            )
        shapes = particle_shapes["shape"].astype(str).tolist()

        sam_image = sam_mask_overlay_notext(processed, masks)
        shape_image = shape_mask_overlay_notext(
            processed,
            masks,
            shapes,
            color_map,
        )
        (
            _,
            boundary_pixels,
            analysis_centroids,
            concave_vertex,
            voronoi_areas,
            _,
        ) = spatial_geometry_from_masks(masks)
        voronoi_figure, voronoi_image = create_voronoi_overlay(
            processed,
            boundary_pixels,
            voronoi_areas,
            scale=None,
            figsize=(10, 10),
            concave_vertex=concave_vertex,
            inside_centroids=analysis_centroids,
            show_area_labels=False,
            show_title=False,
            colorbar_label="Relative Voronoi Cell Area",
            relative_colorbar=True,
        )
        plt.close(voronoi_figure)

        expected_shape = processed.shape[:2]
        output_images = {
            "sam": sam_image,
            "shape": shape_image,
            "voronoi": voronoi_image,
        }
        for asset_name, image in output_images.items():
            if image.shape[:2] != expected_shape:
                raise AssertionError(
                    f"{spec.image_id} {asset_name} asset changed image dimensions: "
                    f"{image.shape[:2]} vs {expected_shape}"
                )

        write_rgb_image(
            destination / FIGURE_ASSET_FILENAMES["sam"], sam_image
        )
        write_rgb_image(
            destination / FIGURE_ASSET_FILENAMES["shape"], shape_image
        )
        write_bgr_image(
            destination / FIGURE_ASSET_FILENAMES["voronoi"], voronoi_image
        )

        record = {
            "image_id": spec.image_id,
            "pixel_height": int(expected_shape[0]),
            "pixel_width": int(expected_shape[1]),
            "mask_count": len(masks),
            "text_annotations": False,
            "sam_color": rgb_to_hex(SAM_MASK_COLOR_RGB),
            "shape_colors": {
                label: rgb_to_hex(color_map[label]) for label in GT_SHAPE_LABELS
            },
            "sources": {
                "preprocessed_image": str(preprocessed_path.resolve()),
                "masks": str(masks_path.resolve()),
                "shape_predictions": str(shapes_path.resolve()),
            },
            "outputs": {
                name: str((destination / filename).resolve())
                for name, filename in FIGURE_ASSET_FILENAMES.items()
            },
        }
        write_json(destination / "figure_assets.json", record)
        records.append(record)
        print(
            f"  [OK] Figure assets {spec.image_id}: "
            + ", ".join(FIGURE_ASSET_FILENAMES.values())
        )

    expected_ids = [spec.image_id for spec in plan if spec.group == "gt_evaluation"]
    exported_ids = [record["image_id"] for record in records]
    if exported_ids != expected_ids:
        raise AssertionError(
            f"Figure asset export set differs from the GT plan: {exported_ids}"
        )

    manifest_path = output_dir / "gt_evaluation" / "figure_assets_manifest.json"
    write_json(
        manifest_path,
        {
            "created_at": now_string(),
            "source": "completed_gt_analysis_caches",
            "model_inference_rerun": False,
            "image_count": len(records),
            "assets_per_image": list(FIGURE_ASSET_FILENAMES.values()),
            "images": records,
        },
    )
    print(f"[OK] Saved figure asset manifest: {manifest_path}")
    return manifest_path


def error_record(spec: InputSpec, stage: str, error: Exception) -> Dict[str, Any]:
    return {
        "image_id": spec.image_id,
        "group": spec.group,
        "source_image": str(spec.image_path),
        "stage": stage,
        "error": str(error),
    }


def save_error(destination: Path, stage: str, error: Exception) -> None:
    write_json(
        destination / f"{stage}_error.json",
        {"status": "error", "stage": stage, "created_at": now_string(), "error": str(error)},
    )


def clear_error(destination: Path, stage: str) -> None:
    try:
        (destination / f"{stage}_error.json").unlink()
    except FileNotFoundError:
        pass


def build_gt_distribution_plot_frame(values: Sequence[float]) -> pd.DataFrame:
    """Build Origin coordinates from the same calculations used by the PNG plot."""
    plot_data = spatial_distribution_plot_data(values)
    if not plot_data:
        raise ValueError("Cannot build spatial plot coordinates without Voronoi areas")

    summary = plot_data["summary"]
    kde_x = np.asarray(plot_data["kde_x"], dtype=np.float64)
    kde_y_normalized = np.asarray(
        plot_data["kde_y_normalized"], dtype=np.float64
    )
    observed_y = np.asarray(plot_data["observed_y"], dtype=np.float64)
    # Keep the deterministic jitter while centering points between IQR and KDE.
    observed_y = observed_y - observed_y.mean() + (0.55 + 0.15) / 2.0

    def series(data: Sequence[float]) -> pd.Series:
        return pd.Series(np.asarray(data, dtype=np.float64), dtype="float64")

    frame = pd.DataFrame(
        {
            "Observed_X": series(plot_data["observed_x"]),
            "Observed_Y": series(observed_y),
            "KDE_X": series(kde_x),
            "KDE_Y_Normalized": series(kde_y_normalized),
            "KDE_Y_Plot": series(0.55 + 0.40 * kde_y_normalized),
            "KDE_Base_Y": series(np.full(kde_x.shape, 0.55, dtype=np.float64)),
            "IQR_X": series([summary["q1"], summary["q3"]]),
            "IQR_Y": series([0.15, 0.15]),
            "Median_CI_X": series(
                [summary["median_ci_low"], summary["median_ci_high"]]
            ),
            "Median_CI_Y": series([0.05, 0.05]),
            "Median_X": series([summary["median"]]),
            "Median_Y": series([0.05]),
        },
        columns=GT_DIST_PLOT_COLUMNS,
    )

    if len(frame) != len(kde_x):
        raise AssertionError("Origin spatial plot frame must have one row per KDE X value")
    if float(frame["KDE_Y_Normalized"].max()) != 1.0:
        raise AssertionError("Normalized KDE maximum is not 1")
    observed_y_values = frame["Observed_Y"].dropna()
    if not ((observed_y_values > 0.15) & (observed_y_values < 0.55)).all():
        raise AssertionError("Origin observed-value Y coordinates must lie between IQR and KDE")
    expected_plot_y = 0.55 + 0.40 * frame["KDE_Y_Normalized"]
    if not np.array_equal(frame["KDE_Y_Plot"].to_numpy(), expected_plot_y.to_numpy()):
        raise AssertionError("Origin KDE plot Y values do not match the requested transform")
    return frame


def export_gt_origin_plot_workbook(
    plan: Sequence[InputSpec], output_dir: Path
) -> Path:
    """Export Origin-ready raw GT-group plot data to one workbook."""
    gt_specs = {
        spec.image_id: spec
        for spec in plan
        if spec.group == "gt_evaluation"
    }
    missing_specs = [image_id for image_id in GT_IMAGE_STEMS if image_id not in gt_specs]
    if missing_specs:
        raise ValueError(
            "GT Origin export requires all four image specs; missing: "
            + ", ".join(missing_specs)
        )

    sheets: Dict[str, pd.DataFrame] = {}
    all_size_frames: List[pd.DataFrame] = []
    shape_counts_by_image: Dict[str, Dict[str, int]] = {}

    for image_id in GT_IMAGE_STEMS:
        spec = gt_specs[image_id]
        destination = image_dir(output_dir, spec)
        particle_path = destination / "particle_shapes.csv"
        distribution_path = destination / "spatial_voronoi_areas.csv"
        missing_sources = [
            str(path)
            for path in (particle_path, distribution_path)
            if not path.exists()
        ]
        if missing_sources:
            raise FileNotFoundError(
                f"GT Origin source data are incomplete for {image_id}: "
                + ", ".join(missing_sources)
            )

        particles = pd.read_csv(particle_path)
        required_particle_columns = {"particle_id", "area_px", "shape"}
        missing_particle_columns = required_particle_columns - set(particles.columns)
        if missing_particle_columns:
            raise ValueError(
                f"{particle_path} is missing columns: "
                + ", ".join(sorted(missing_particle_columns))
            )
        if particles.empty:
            raise ValueError(f"No particle rows found for GT image {image_id}")
        if particles["particle_id"].duplicated().any():
            raise ValueError(f"Duplicate particle IDs found for GT image {image_id}")

        particle_ids = pd.to_numeric(particles["particle_id"], errors="raise")
        area_values = pd.to_numeric(particles["area_px"], errors="raise")
        if not np.isfinite(area_values.to_numpy(dtype=np.float64)).all():
            raise ValueError(f"Non-finite particle area found for GT image {image_id}")
        if (area_values < 0).any():
            raise ValueError(f"Negative particle area found for GT image {image_id}")

        shape_values = particles["shape"].astype("string")
        if shape_values.isna().any():
            raise ValueError(f"Missing shape label found for GT image {image_id}")
        unexpected_shapes = sorted(set(shape_values) - set(GT_SHAPE_LABELS))
        if unexpected_shapes:
            raise ValueError(
                f"Unexpected shape labels for GT image {image_id}: {unexpected_shapes}"
            )

        size_df = pd.DataFrame(
            {
                "Image_ID": [str(image_id)] * len(particles),
                "Particle_ID": particle_ids.to_numpy(),
                "Area_px2": area_values.to_numpy(),
            }
        )
        all_size_frames.append(size_df.copy())
        sheets[f"{image_id}_size_hist"] = size_df

        counts = {
            label: int((shape_values == label).sum())
            for label in GT_SHAPE_LABELS
        }
        shape_counts_by_image[image_id] = counts
        total_shapes = sum(counts.values())
        if total_shapes != len(particles):
            raise AssertionError(
                f"Shape counts do not equal particle count for GT image {image_id}"
            )
        shape_df = pd.DataFrame(
            {
                "Image_ID": [str(image_id)] * len(GT_SHAPE_LABELS),
                "Shape": list(GT_SHAPE_LABELS),
                "Count": [counts[label] for label in GT_SHAPE_LABELS],
                "Percent": [
                    counts[label] / total_shapes * 100.0
                    if total_shapes
                    else 0.0
                    for label in GT_SHAPE_LABELS
                ],
            }
        )
        sheets[f"{image_id}_shape_pie"] = shape_df

        distribution = pd.read_csv(distribution_path, float_precision="round_trip")
        if "voronoi_area_px2" not in distribution.columns:
            raise ValueError(
                f"{distribution_path} is missing column: voronoi_area_px2"
            )
        voronoi_values = pd.to_numeric(
            distribution["voronoi_area_px2"], errors="raise"
        )
        if distribution.empty:
            raise ValueError(f"No Voronoi rows found for GT image {image_id}")
        if not np.isfinite(voronoi_values.to_numpy(dtype=np.float64)).all():
            raise ValueError(f"Non-finite Voronoi area found for GT image {image_id}")
        if (voronoi_values < 0).any():
            raise ValueError(f"Negative Voronoi area found for GT image {image_id}")

        distribution_df = pd.DataFrame(
            {
                "Image_ID": [str(image_id)] * len(distribution),
                "Voronoi_Cell_ID": np.arange(1, len(distribution) + 1, dtype=int),
                "Voronoi_Area_px2": voronoi_values.to_numpy(),
            }
        )
        sheets[f"{image_id}_dist_violin"] = distribution_df

        sheets[f"{image_id}_dist_plot"] = build_gt_distribution_plot_frame(
            voronoi_values.to_numpy(dtype=np.float64)
        )

    all_size_df = pd.concat(all_size_frames, ignore_index=True)
    sheets["all_size_hist"] = all_size_df

    all_shape_counts = {
        label: sum(
            shape_counts_by_image[image_id][label]
            for image_id in GT_IMAGE_STEMS
        )
        for label in GT_SHAPE_LABELS
    }
    all_shape_total = sum(all_shape_counts.values())
    if all_shape_total != len(all_size_df):
        raise AssertionError(
            "Combined GT shape counts do not equal combined particle count"
        )
    all_shape_df = pd.DataFrame(
        {
            "Shape": list(GT_SHAPE_LABELS),
            "Count": [all_shape_counts[label] for label in GT_SHAPE_LABELS],
            "Percent": [
                all_shape_counts[label] / all_shape_total * 100.0
                if all_shape_total
                else 0.0
                for label in GT_SHAPE_LABELS
            ],
        }
    )
    sheets["all_shape_pie"] = all_shape_df

    metadata_df = pd.DataFrame(
        [
            {"Item": "Number_of_images", "Value": len(GT_IMAGE_STEMS)},
            {"Item": "Image_IDs", "Value": ";".join(GT_IMAGE_STEMS)},
            {
                "Item": "Analysis_population",
                "Value": "Final SAM masks for the GT-evaluation image group",
            },
            {
                "Item": "Preprocessing",
                "Value": "BM3D + Noise2SR",
            },
            {
                "Item": "Size_source",
                "Value": "particle_shapes.csv area_px; raw mask pixel area",
            },
            {
                "Item": "Shape_source",
                "Value": "particle_shapes.csv shape; CLIP classification",
            },
            {
                "Item": "Distribution_source",
                "Value": (
                    "spatial_voronoi_areas.csv voronoi_area_px2; "
                    "VISION convex-hull-clipped Voronoi cells"
                ),
            },
            {
                "Item": "Spatial_KDE_settings",
                "Value": (
                    f"Gaussian; grid_points={SPATIAL_KDE_GRID_POINTS}; "
                    f"x_padding_fraction={SPATIAL_KDE_X_PADDING_FRACTION}; "
                    f"GUI_height={SPATIAL_KDE_HEIGHT}"
                ),
            },
            {
                "Item": "Spatial_median_CI",
                "Value": (
                    "95% percentile bootstrap; "
                    f"resamples={SPATIAL_BOOTSTRAP_RESAMPLES}; "
                    f"seed={SPATIAL_BOOTSTRAP_SEED}"
                ),
            },
            {
                "Item": "Combined_outputs",
                "Value": "Size histogram and shape pie only, as requested",
            },
            {"Item": "Created_at", "Value": now_string()},
        ],
        columns=["Item", "Value"],
    )
    sheets["metadata"] = metadata_df

    workbook_path = output_dir / GT_ORIGIN_WORKBOOK_NAME
    workbook_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
        for sheet_name, dataframe in sheets.items():
            dataframe.to_excel(writer, sheet_name=sheet_name, index=False)

    from openpyxl import load_workbook

    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != list(sheets):
            raise AssertionError(
                f"Unexpected GT Origin workbook sheets: {workbook.sheetnames}"
            )
        for sheet_name, dataframe in sheets.items():
            worksheet = workbook[sheet_name]
            headers = [
                cell.value
                for cell in next(worksheet.iter_rows(min_row=1, max_row=1))
            ]
            if headers != list(dataframe.columns):
                raise AssertionError(
                    f"Unexpected columns in {sheet_name}: {headers}"
                )
            if worksheet.max_row - 1 != len(dataframe):
                raise AssertionError(
                    f"Unexpected row count in {sheet_name}: {worksheet.max_row - 1}"
                )
            for row in worksheet.iter_rows():
                if any(cell.data_type == "f" for cell in row):
                    raise AssertionError(f"Formula cell found in {sheet_name}")

        id_sheets = [
            name
            for name, dataframe in sheets.items()
            if len(dataframe.columns) > 0 and dataframe.columns[0] == "Image_ID"
        ]
        for sheet_name in id_sheets:
            worksheet = workbook[sheet_name]
            for row in worksheet.iter_rows(
                min_row=2,
                min_col=1,
                max_col=1,
            ):
                cell = row[0]
                if cell.data_type != "s":
                    raise AssertionError(
                        f"Image ID is not text in {sheet_name}: {cell.value}"
                    )
    finally:
        workbook.close()

    print("\n--- GT Origin Workbook ---")
    for sheet_name, dataframe in sheets.items():
        print(f"  {sheet_name}: {len(dataframe)} data rows")
    print(f"  Combined particles: {len(all_size_df)}")
    print(f"  Combined shape counts: {all_shape_counts}")
    print(f"[OK] Saved GT Origin workbook: {workbook_path}")
    return workbook_path


def export_batch_summaries(
    plan: Sequence[InputSpec], output_dir: Path, errors: Sequence[Mapping[str, Any]]
) -> None:
    all_rows: List[Dict[str, Any]] = []
    gt_rows: List[Dict[str, Any]] = []
    gt_spatial_rows: List[Dict[str, Any]] = []
    gt_spatial_area_rows: List[pd.DataFrame] = []
    gt_shape_rows: List[Dict[str, Any]] = []
    gt_shape_particle_rows: List[pd.DataFrame] = []

    for spec in plan:
        destination = image_dir(output_dir, spec)
        segmentation = read_json(segmentation_summary_path(destination)) or {}
        spatial = read_json(spatial_summary_path(destination)) or {}
        shape = read_json(shape_summary_path(destination)) or {}
        segmentation_info = segmentation.get("segmentation", {})
        row = {
            "image_id": spec.image_id,
            "group": spec.group,
            "source_image": str(spec.image_path),
            "segmentation_status": segmentation.get("status", "not_completed"),
            "raw_mask_count": segmentation_info.get("raw_mask_count"),
            "background_removed_count": segmentation_info.get("background_removed_count"),
            "final_mask_count": segmentation_info.get("final_mask_count"),
            "processed_shape": "x".join(str(value) for value in segmentation.get("processed_shape", [])),
            "shape_status": shape.get("status", "not_completed"),
            "spatial_status": spatial.get("status", "not_completed"),
            "figure_assets_complete": all(
                (destination / filename).exists()
                for filename in FIGURE_ASSET_FILENAMES.values()
            ),
        }
        all_rows.append(row)

        evaluation = segmentation.get("gt_evaluation")
        if evaluation:
            gt_rows.append({"image_id": spec.image_id, **evaluation})

        metrics = spatial.get("spatial_metrics", {})
        clustering = spatial.get("clustering_metrics", {})
        gt_spatial_rows.append(
            {
                "image_id": spec.image_id,
                "source_image": str(spec.image_path),
                "spatial_status": spatial.get("status", "not_completed"),
                "measurement_unit": spatial.get("measurement_unit", "px"),
                **metrics,
                **clustering,
            }
        )
        area_path = destination / "spatial_voronoi_areas.csv"
        if area_path.exists():
            areas = pd.read_csv(area_path)
            areas.insert(0, "image_id", spec.image_id)
            gt_spatial_area_rows.append(areas)

        metrics = shape.get("shape_metrics", {})
        counts = shape.get("shape_counts_all_labels", {})
        gt_shape_rows.append(
            {
                "image_id": spec.image_id,
                "source_image": str(spec.image_path),
                "shape_status": shape.get("status", "not_completed"),
                "particle_count": shape.get("particle_count"),
                "dominant_shape": metrics.get("dominant_shape"),
                "dominant_percentage": metrics.get("dominant_percentage"),
                "avg_confidence": metrics.get("avg_confidence"),
                "shannon_entropy": metrics.get("shannon_entropy"),
                "simpson_diversity": metrics.get("simpson_diversity"),
                **{label: counts.get(label, 0) for label in GT_SHAPE_LABELS},
            }
        )
        particle_path = destination / "particle_shapes.csv"
        if particle_path.exists():
            particles = pd.read_csv(particle_path)
            particles.insert(0, "image_id", spec.image_id)
            gt_shape_particle_rows.append(particles)

    all_df = pd.DataFrame(all_rows)
    gt_df = pd.DataFrame(gt_rows)
    gt_spatial_df = pd.DataFrame(gt_spatial_rows)
    gt_spatial_areas_df = (
        pd.concat(gt_spatial_area_rows, ignore_index=True)
        if gt_spatial_area_rows
        else pd.DataFrame()
    )
    gt_shape_df = pd.DataFrame(gt_shape_rows)
    gt_shape_particles_df = (
        pd.concat(gt_shape_particle_rows, ignore_index=True)
        if gt_shape_particle_rows
        else pd.DataFrame()
    )
    errors_df = pd.DataFrame(list(errors))

    all_df.to_csv(output_dir / "all_images_summary.csv", index=False)
    gt_df.to_csv(output_dir / "segmentation_gt_metrics.csv", index=False)
    gt_spatial_df.to_csv(output_dir / "gt_spatial_uniformity_metrics.csv", index=False)
    gt_spatial_areas_df.to_csv(output_dir / "gt_spatial_voronoi_areas.csv", index=False)
    gt_shape_df.to_csv(output_dir / "gt_shape_summary.csv", index=False)
    gt_shape_particles_df.to_csv(output_dir / "gt_particle_shapes.csv", index=False)
    errors_df.to_csv(output_dir / "errors.csv", index=False)

    workbook_path = output_dir / "case_study_summary.xlsx"
    try:
        with pd.ExcelWriter(workbook_path, engine="openpyxl") as writer:
            all_df.to_excel(writer, sheet_name="all_images", index=False)
            gt_df.to_excel(writer, sheet_name="gt_segmentation", index=False)
            gt_spatial_df.to_excel(writer, sheet_name="gt_spatial", index=False)
            gt_spatial_areas_df.to_excel(writer, sheet_name="gt_voronoi_areas", index=False)
            gt_shape_df.to_excel(writer, sheet_name="gt_shape_summary", index=False)
            gt_shape_particles_df.to_excel(writer, sheet_name="gt_particle_shapes", index=False)
            errors_df.to_excel(writer, sheet_name="errors", index=False)
    except Exception as exc:
        print(f"[WARN] Could not create Excel workbook: {exc}")

    try:
        export_gt_origin_plot_workbook(plan, output_dir)
    except Exception as exc:
        print(f"[WARN] Could not create GT Origin workbook: {exc}")


def print_plan(plan: Sequence[InputSpec], output_dir: Path) -> None:
    print("GT Case Study batch plan")
    print(f"  Output: {output_dir}")
    print("  Segmentation + shape + spatial uniformity:")
    for spec in plan:
        print(f"    - {spec.image_path.name} vs {spec.gt_path.name}")


def main() -> int:
    args = parse_args()
    args.case_study_dir = args.case_study_dir.resolve()
    args.output_dir = args.output_dir.resolve()
    full_plan = build_plan(args.case_study_dir)
    plan = full_plan
    missing_files = validate_plan(plan)
    print_plan(plan, args.output_dir)
    if missing_files:
        print("  [WARN] Missing inputs will be skipped while available images continue:")
        for missing_path in missing_files:
            print(f"    - {missing_path}")

    if args.dry_run:
        print("Dry run complete. No preprocessing or model inference was started.")
        return 1 if missing_files else 0

    args.output_dir.mkdir(parents=True, exist_ok=True)
    if args.figure_assets_only:
        try:
            export_gt_figure_assets(plan, args.output_dir)
        except Exception as exc:
            print(f"[ERROR] Text-free figure asset export failed: {exc}")
            return 1
        return 0
    if args.gt_origin_export_only:
        try:
            export_gt_origin_plot_workbook(plan, args.output_dir)
        except Exception as exc:
            print(f"[ERROR] GT Origin workbook export failed: {exc}")
            return 1
        return 0

    preproc_hash = stable_hash(preprocessing_config(args))
    segmentation_hashes = {
        spec.image_id: stable_hash(
            {
                **segmentation_config(VISION_SAM_CONFIG),
                "preprocessing_hash": preproc_hash,
                "gt_evaluation": {
                    "match_iou": args.match_iou,
                    "gt_min_area": args.gt_min_area,
                },
            }
        )
        for spec in full_plan
    }
    spatial_hashes = {
        spec.image_id: stable_hash(
            {
                **spatial_config(args),
                "segmentation_hash": segmentation_hashes[spec.image_id],
            }
        )
        for spec in full_plan
        if spec.group == "gt_evaluation"
    }
    default_seg_hash = stable_hash(
        {
            **segmentation_config(VISION_SAM_CONFIG),
            "preprocessing_hash": preproc_hash,
            "gt_evaluation": {
                "match_iou": args.match_iou,
                "gt_min_area": args.gt_min_area,
            },
        }
    )
    shape_hashes = {
        spec.image_id: stable_hash(
            {
                **shape_config(args),
                "segmentation_hash": segmentation_hashes[spec.image_id],
            }
        )
        for spec in full_plan
        if spec.group == "gt_evaluation"
    }
    manifest = {
        "status": "running",
        "pipeline_version": PIPELINE_VERSION,
        "started_at": now_string(),
        "case_study_dir": str(args.case_study_dir),
        "output_dir": str(args.output_dir),
        "preprocessing_config": preprocessing_config(args),
        "sam_config": VISION_SAM_CONFIG,
        "spatial_config": spatial_config(args),
        "shape_config": shape_config(args),
        "hashes": {
            "preprocessing": preproc_hash,
            "segmentation_default": default_seg_hash,
            "spatial": spatial_hashes,
            "shape": shape_hashes,
        },
        "inputs": [
            {
                "image_id": spec.image_id,
                "group": spec.group,
                "image": str(spec.image_path),
                "ground_truth": str(spec.gt_path) if spec.gt_path else None,
            }
            for spec in plan
        ],
    }
    write_json(args.output_dir / "run_manifest.json", manifest)

    errors: List[Dict[str, Any]] = []
    preprocessing_results: Dict[str, Tuple[Dict[str, Any], Path, bool]] = {}
    preprocessing_ran_ids = set()
    segmentation_ran_ids = set()
    preproc_cache_dir = args.output_dir / "preprocessed_bm3d_noise2sr"

    print("\n" + "=" * 80)
    print("PHASE 1/4: BM3D + Noise2SR PREPROCESSING")
    print("=" * 80)
    print(f"Reusable preprocessing cache: {preproc_cache_dir}")
    print("Completed images are reused on the next normal run.")
    for spec in plan:
        try:
            _, metadata, preprocessed_path, reused = load_or_preprocess(
                spec, args, preproc_hash
            )
            preprocessing_results[spec.image_id] = (metadata, preprocessed_path, reused)
            clear_error(image_dir(args.output_dir, spec), "preprocessing")
            if not reused:
                preprocessing_ran_ids.add(spec.image_id)
        except Exception as exc:
            errors.append(error_record(spec, "preprocessing", exc))
            save_error(image_dir(args.output_dir, spec), "preprocessing", exc)
            print(f"  [ERROR] Preprocessing {spec.image_id}: {exc}")
            cleanup_gpu()

    segmentation_pending = [
        spec
        for spec in plan
        if spec.image_id in preprocessing_results
        and (
            spec.image_id in preprocessing_ran_ids
            or not segmentation_cache_valid(
                spec, args, segmentation_hashes[spec.image_id]
            )
        )
    ]
    print("\n" + "=" * 80)
    print("PHASE 2/4: SAM SEGMENTATION + BACKGROUND REMOVAL")
    print("=" * 80)
    if segmentation_pending:
        try:
            import torch

            device = "cuda" if torch.cuda.is_available() else "cpu"
            print(f"Loading SAM2 Hiera-L once on {device}...")
            sam_model = load_sam_model("hiera_l", device=device)
            for spec in segmentation_pending:
                metadata, preprocessed_path, _ = preprocessing_results[spec.image_id]
                try:
                    print(f"  [INFO] {spec.image_id}: SAM mode=default")
                    run_segmentation_for_image(
                        sam_model,
                        spec,
                        args,
                        metadata,
                        preprocessed_path,
                        segmentation_hashes[spec.image_id],
                        VISION_SAM_CONFIG,
                    )
                    clear_error(image_dir(args.output_dir, spec), "segmentation")
                    clear_error(image_dir(args.output_dir, spec), "sam_initialization")
                    segmentation_ran_ids.add(spec.image_id)
                except Exception as exc:
                    errors.append(error_record(spec, "segmentation", exc))
                    save_error(image_dir(args.output_dir, spec), "segmentation", exc)
                    print(f"  [ERROR] Segmentation {spec.image_id}: {exc}")
                    cleanup_gpu()
            del sam_model
            cleanup_gpu()
        except Exception as exc:
            print(f"  [ERROR] Unable to initialize SAM2: {exc}")
            for spec in segmentation_pending:
                errors.append(error_record(spec, "sam_initialization", exc))
                save_error(image_dir(args.output_dir, spec), "sam_initialization", exc)
    else:
        print("  [REUSE] All segmentation results are already complete.")

    spatial_targets = [
        spec
        for spec in plan
        if spec.image_id in preprocessing_results
        and segmentation_artifacts_valid(
            spec, args.output_dir, segmentation_hashes[spec.image_id]
        )
        and (
            spec.image_id in segmentation_ran_ids
            or not spatial_cache_valid(spec, args, spatial_hashes[spec.image_id])
        )
    ]
    print("\n" + "=" * 80)
    print("PHASE 3/4: SPATIAL UNIFORMITY")
    print("=" * 80)
    if spatial_targets:
        for spec in spatial_targets:
            _, preprocessed_path, _ = preprocessing_results[spec.image_id]
            try:
                run_spatial_for_image(
                    spec, args, preprocessed_path, spatial_hashes[spec.image_id]
                )
                clear_error(image_dir(args.output_dir, spec), "spatial_uniformity")
            except Exception as exc:
                errors.append(error_record(spec, "spatial_uniformity", exc))
                save_error(image_dir(args.output_dir, spec), "spatial_uniformity", exc)
                print(f"  [ERROR] Spatial uniformity {spec.image_id}: {exc}")
    else:
        print("  [REUSE] All requested spatial uniformity results are already complete.")

    shape_targets = [
        spec
        for spec in plan
        if segmentation_artifacts_valid(
            spec, args.output_dir, segmentation_hashes[spec.image_id]
        )
        and (
            spec.image_id in segmentation_ran_ids
            or not shape_cache_valid(spec, args, shape_hashes[spec.image_id])
        )
    ]
    print("\n" + "=" * 80)
    print("PHASE 4/4: CLIP SHAPE CLASSIFICATION")
    print("=" * 80)
    if shape_targets:
        try:
            clip_resources: Optional[Mapping[str, Any]] = None
            targets_with_masks = []
            for spec in shape_targets:
                masks = load_masks(image_dir(args.output_dir, spec) / "predicted_masks.npz")
                if masks:
                    targets_with_masks.append(spec)
                else:
                    try:
                        run_shape_for_image(
                            spec, args, None, shape_hashes[spec.image_id]
                        )
                        clear_error(
                            image_dir(args.output_dir, spec), "shape_classification"
                        )
                    except Exception as exc:
                        errors.append(error_record(spec, "shape_classification", exc))
                        save_error(image_dir(args.output_dir, spec), "shape_classification", exc)

            if targets_with_masks:
                clip_resources = load_clip_resources(args.clip_model)
                for spec in targets_with_masks:
                    try:
                        run_shape_for_image(
                            spec, args, clip_resources, shape_hashes[spec.image_id]
                        )
                        clear_error(
                            image_dir(args.output_dir, spec), "shape_classification"
                        )
                        clear_error(image_dir(args.output_dir, spec), "clip_initialization")
                    except Exception as exc:
                        errors.append(error_record(spec, "shape_classification", exc))
                        save_error(image_dir(args.output_dir, spec), "shape_classification", exc)
                        print(f"  [ERROR] Shape classification {spec.image_id}: {exc}")
                        cleanup_gpu()
                if clip_resources is not None:
                    del clip_resources
                cleanup_gpu()
        except Exception as exc:
            print(f"  [ERROR] Unable to initialize CLIP: {exc}")
            for spec in shape_targets:
                errors.append(error_record(spec, "clip_initialization", exc))
                save_error(image_dir(args.output_dir, spec), "clip_initialization", exc)
    else:
        print("  [REUSE] All requested shape results are already complete or await segmentation.")

    print("\n" + "=" * 80)
    print("TEXT-FREE FIGURE ASSETS")
    print("=" * 80)
    try:
        export_gt_figure_assets(plan, args.output_dir)
    except Exception as exc:
        print(f"  [ERROR] Text-free figure asset export failed: {exc}")
        errors.append(
            {
                "image_id": "gt_evaluation",
                "group": "gt_evaluation",
                "source_image": None,
                "stage": "figure_asset_export",
                "error": str(exc),
            }
        )

    export_batch_summaries(full_plan, args.output_dir, errors)
    manifest["status"] = "completed" if not errors else "completed_with_errors"
    manifest["finished_at"] = now_string()
    manifest["errors"] = errors
    write_json(args.output_dir / "run_manifest.json", manifest)

    print("\n" + "=" * 80)
    print("CASE STUDY BATCH COMPLETE")
    print("=" * 80)
    print(f"Output: {args.output_dir}")
    print(
        "Summary CSV: segmentation_gt_metrics.csv, "
        "gt_spatial_uniformity_metrics.csv, gt_shape_summary.csv, "
        "gt_particle_shapes.csv"
    )
    print("Summary workbook: case_study_summary.xlsx")
    print(f"GT Origin workbook: {GT_ORIGIN_WORKBOOK_NAME}")
    print(
        "Text-free assets per case: "
        + ", ".join(FIGURE_ASSET_FILENAMES.values())
    )
    if errors:
        print(f"Completed with {len(errors)} error(s). See: {args.output_dir / 'errors.csv'}")
        return 1
    print("Completed successfully.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
